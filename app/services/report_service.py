# -*- coding: utf-8 -*-
"""报告生成：生成 Word 评价报告。

格式（论文规范）：
- 题目：三号(16pt) 黑体 居中
- 一级标题：四号(14pt) 黑体 居中
- 正文/表格：小四号(12pt) 宋体
- 行距：单倍行距
- 全部黑色
"""
import io
import json

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

from app.utils.timefmt import to_local, now_utc
from app.services import workflow as wf

BLACK = RGBColor(0, 0, 0)


def _set_run(run, text=None, size=12, bold=False, east_asia="宋体"):
    if text is not None:
        run.text = text
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    return run


def _set_line_spacing(p):
    """单倍行距。"""
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _para(doc, text, size=12, bold=False, align=None, space_after=6, east_asia="宋体"):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    _set_line_spacing(p)
    _set_run(p.add_run(), text, size, bold, east_asia)
    return p


def _heading(doc, text, level=1):
    """level=0 题目（三号黑体居中）；level=1 一级标题（四号黑体居中）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    _set_line_spacing(p)
    if level == 0:
        size = 16  # 三号
    else:
        size = 14  # 四号
    _set_run(p.add_run(), text, size=size, bold=True, east_asia="黑体")
    return p


def _table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        cell_p = hdr[i].paragraphs[0]
        _set_line_spacing(cell_p)
        _set_run(cell_p.add_run(), htext, 12, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell_p = cells[i].paragraphs[0]
            _set_line_spacing(cell_p)
            _set_run(cell_p.add_run(), str(val), 12)
    return t


def build_report(patch, mine_result, record=None, weak=None, operator=None):
    """生成评价报告 Word 文档，返回 BytesIO。"""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 题目：三号黑体居中
    _heading(doc, "历史遗留废弃矿山自然恢复适宜性评价报告", 0)
    _para(doc, "软件：历史遗留废弃矿山自然恢复适宜性评价系统 V1.0",
          size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "生成时间：%s" % to_local(now_utc(), "%Y-%m-%d %H:%M:%S"),
          size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # 计算人与参数来源
    if record is not None:
        from app.views.evaluate import param_source_summary
        try:
            params = json.loads(record.params_json) if record.params_json else {}
        except ValueError:
            params = {}
        meta = [
            ("计算人", operator or "—"),
            ("计算时间", to_local(record.created_at, "%Y-%m-%d %H:%M:%S")),
            ("评价记录编号", record.id),
            ("计算参数", param_source_summary(params)),
        ]
        _table(doc, ["项目", "内容"], meta)

    # 图斑信息
    _heading(doc, "一、图斑基本信息", 1)
    info = [
        ("图斑名称", patch.name),
        ("所在行政区", patch.admin_region or "—"),
        ("图斑面积（公顷）", str(patch.area_ha) if patch.area_ha is not None else "—"),
        ("开采矿种", patch.mineral_type or "—"),
        ("开采方式", patch.mining_method or "—"),
    ]
    _table(doc, ["项目", "内容"], info)

    # 评价结论
    _heading(doc, "二、评价结论", 1)
    gi = mine_result["grade_info"]
    concl = [
        ("最终评价因子 F", "%.4f" % mine_result["final_score"]),
        ("适宜性等级", mine_result["grade"]),
        ("判定依据", gi["meaning"]),
        ("推荐措施", gi["measure"]),
    ]
    _table(doc, ["项目", "结果"], concl)

    # 各单元得分
    _heading(doc, "三、各微地质单元得分", 1)
    rows = []
    for u in mine_result["unit_results"]:
        rows.append((u["unit_type"], "%.2f%%" % (u["ratio"] * 100),
                     "%.4f" % u["unit_score"], "%.4f" % u["weighted"]))
    _table(doc, ["微地质单元", "面积占比", "单元评价因子 F′", "加权贡献"], rows)

    # 完整得分明细
    _heading(doc, "四、得分明细", 1)
    for u in mine_result["unit_results"]:
        _para(doc, "%s（占比 %.2f%%）" % (u["unit_type"], u["ratio"] * 100),
              size=12, bold=True, space_after=4)
        detail_rows = []
        for d in u["details"]:
            raw = "—" if d["raw"] is None else d["raw"]
            std = "—" if d["std_value"] is None else "%.4f" % d["std_value"]
            w = "%.4f" % d["weight"]
            c = "—" if d["contribution"] is None else "%.4f" % d["contribution"]
            detail_rows.append((d["l1"], d["l2"], raw, std, w, c))
        _table(doc, ["一级指标", "二级指标", "实测值", "标准化值", "权重", "贡献"],
               detail_rows)
        l1_row = []
        for k, v in u["l1_scores"].items():
            l1_row.append("%s=%.4f" % (k, v))
        _para(doc, "一级指标得分：" + "，".join(l1_row), size=12, space_after=10)

    # 数据完整度
    undetected = wf.undetected_metals(patch.id)
    if undetected:
        _heading(doc, "五、数据完整度说明", 1)
        parts = ["%s：%s" % (t, "、".join(m)) for t, m in undetected]
        _para(doc, "以下重金属项未检测，按论文注记以「未受污染」计入污染指标"
                   "（并非实测达标）：%s。" % "；".join(parts), size=12)

    # 短板提示
    if weak:
        _heading(doc, "六、短板提示", 1)
        _para(doc, "按「权重 ×（1 − 标准化值）」排序，提升下列指标对总分的拉动最大。",
              size=12, space_after=4)
        _table(doc, ["微地质单元", "指标", "实测值", "标准化值", "提升空间"],
               [(u, l2, "—" if raw is None else raw, "%.4f" % std,
                 "%.4f" % gain) for u, l2, raw, std, gain in weak])

    # 签字栏
    _para(doc, "", size=12, space_after=18)
    _para(doc, "评价人（签字）：____________________", size=12)
    _para(doc, "日期：____________________", size=12)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf