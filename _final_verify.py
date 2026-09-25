# -*- coding: utf-8 -*-
"""最终核验：用户文件 vs 改色前备份 —— 正文内容必须完全一致，仅颜色改变。"""
import sys, re, zipfile, hashlib
sys.stdout.reconfigure(encoding='utf-8')
import docx

CUR = '软件著作权登记项目计划书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0.docx'
BAK = '软件著作权登记项目计划书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0_改色前备份.docx'

def load(fn):
    z = zipfile.ZipFile(fn)
    return {n: z.read(n) for n in z.namelist()}

cur, bak = load(CUR), load(BAK)

print('=== 1) 正文 document.xml 是否被改动 ===')
hc = hashlib.sha256(cur['word/document.xml']).hexdigest()
hb = hashlib.sha256(bak['word/document.xml']).hexdigest()
print('   改色前:', hb[:40])
print('   当前  :', hc[:40])
print('   结论  :', '✅ 正文零改动' if hc == hb else '❌ 正文被改动了')

print()
print('=== 2) 逐 part 对比（哪些 part 变了）===')
for n in sorted(set(cur) | set(bak)):
    if n not in cur:
        print('  - 删除:', n); continue
    if n not in bak:
        print('  + 新增:', n); continue
    if cur[n] != bak[n]:
        print('  ~ 已修改:', n, '(%d -> %d bytes)' % (len(bak[n]), len(cur[n])))
    # 相同的就不打印

print()
print('=== 3) 段落/表格/文字内容逐项比对 ===')
dc, db = docx.Document(CUR), docx.Document(BAK)
pc = [p.text for p in dc.paragraphs]
pb = [p.text for p in db.paragraphs]
print('   段落数: %d vs %d' % (len(pc), len(pb)))
print('   段落文字完全一致:', '✅' if pc == pb else '❌')
if pc != pb:
    for i, (a, b) in enumerate(zip(pc, pb)):
        if a != b:
            print('     首个差异 @%d:\n       改前: %r\n       当前: %r' % (i, b[:120], a[:120]))
            break

tc = [[c.text for c in r.cells] for t in dc.tables for r in t.rows]
tb = [[c.text for c in r.cells] for t in db.tables for r in t.rows]
print('   表格数: %d vs %d' % (len(dc.tables), len(db.tables)))
print('   表格文字完全一致:', '✅' if tc == tb else '❌')

# 字符总数
import itertools
n_char_c = sum(len(x) for x in pc) + sum(len(c) for row in tc for c in row)
n_char_b = sum(len(x) for x in pb) + sum(len(c) for row in tb for c in row)
print('   总字符数: %d vs %d' % (n_char_c, n_char_b))

print()
print('=== 4) 标题最终颜色与字体 ===')
s = cur['word/styles.xml'].decode('utf-8')
rows = []
for m in re.finditer(r'<w:style\b[^>]*?w:styleId="([^"]+)"[^>]*?>(.*?)</w:style>', s, re.S):
    sid, blk = m.group(1), m.group(2)
    nm = re.search(r'<w:name w:val="([^"]*)"', blk)
    name = nm.group(1) if nm else '?'
    if not re.match(r'heading \d|Title|TOC Heading|Subtitle', name, re.I):
        continue
    col = re.search(r'<w:color\b[^>]*?w:val="([0-9A-Fa-f]{6})"', blk)
    ea = re.search(r'w:eastAsia="([^"]*)"', blk)
    rows.append((name, col.group(1) if col else '(继承)', ea.group(1) if ea else '(继承)'))
for name, col, ea in rows:
    mark = '✅' if col in ('000000', '(继承)') else '❌'
    print('   %s %-14s 颜色=#%-8s 中文字体=%s' % (mark, name, col, ea))

print()
print('=== 5) 主题强调色 ===')
th = cur['word/theme/theme1.xml'].decode('utf-8')
acc = re.findall(r'<a:accent(\d)>\s*<a:srgbClr val="([0-9A-Fa-f]{6})"', th)
print('  accent1-6:', acc)
print('  是否全黑:', '✅' if all(v == '000000' for _, v in acc) else '❌')
