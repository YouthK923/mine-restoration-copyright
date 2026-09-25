# -*- coding: utf-8 -*-
"""M9 报告输出自检：得分明细页 + Word 导出。"""
import os
import sys
import io
sys.stdout.reconfigure(encoding="utf-8")

from app import create_app

LINE = "=" * 66

class TestConfig:
    SECRET_KEY = "test"
    SQLALCHEMY_DATABASE_URI = "sqlite:///" + os.path.join(
        os.path.dirname(__file__), "app", "app.db")
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    WTF_CSRF_ENABLED = False
    TESTING = True

app = create_app(TestConfig)
client = app.test_client()

from app.models import MinePatch
with app.app_context():
    patch = MinePatch.query.first()
    pid = patch.id

print(LINE)
print("测试 1  报告列表")
print(LINE)
r = client.get("/report/")
print("  状态码:", r.status_code)
assert r.status_code == 200

print()
print(LINE)
print("测试 2  得分明细页")
print(LINE)
r = client.get("/report/%d" % pid)
body = r.get_data(as_text=True)
print("  状态码:", r.status_code)
print("  含「得分明细」:", "得分明细" in body)
print("  含「最终评价因子 F」:", "最终评价因子 F" in body)
print("  含二级指标明细（植被覆盖度/坡度/坡向）:", all(n in body for n in ("植被覆盖度", "坡度", "坡向")))
print("  含权重/标准化值列:", "标准化值" in body and "贡献" in body)
assert r.status_code == 200
assert "植被覆盖度" in body and "标准化值" in body

print()
print(LINE)
print("测试 3  Word 报告导出")
print(LINE)
r = client.get("/report/%d/word" % pid)
print("  状态码:", r.status_code)
print("  Content-Type:", r.headers.get("Content-Type", "")[:60])
data = r.data
print("  文件大小:", len(data), "字节")
print("  是否为有效 docx（PK 头）:", data[:2] == b"PK")
assert r.status_code == 200 and data[:2] == b"PK"

# 用 python-docx 打开验证内容
import docx
d = docx.Document(io.BytesIO(data))
text = "\n".join(p.text for p in d.paragraphs)
print("  含报告标题:", "自然恢复适宜性评价报告" in text)
print("  含评价结论章节:", "评价结论" in text)
print("  含得分明细章节:", "得分明细" in text)
print("  段落数:", len(d.paragraphs), " 表格数:", len(d.tables))
assert "自然恢复适宜性评价报告" in text and len(d.tables) >= 3

print()
print(LINE)
print("M9 报告输出全部测试通过 ✓")
print(LINE)
