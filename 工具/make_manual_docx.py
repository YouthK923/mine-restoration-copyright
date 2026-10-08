# -*- coding: utf-8 -*-
"""生成《软件说明书》的 Word 版。

为什么要有 Word 版：正式提交给审查员的是 PDF（线上申请只收 PDF），
Word 版是给项目成员与师兄直接改字、写批注用的。

两版同源：正文都取自 工具/make_manual.py 的 CONTENT 与 toc_entries()，
改一处两版一起变，不会出现「PDF 改了 Word 没改」这种各说各话的情况。

用法：
    python 工具/make_manual_docx.py
输出：
    软著材料/说明书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0.docx
"""
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOOL_DIR = os.path.join(ROOT, "工具")
if TOOL_DIR not in sys.path:
    sys.path.insert(0, TOOL_DIR)

import make_manual as M          # noqa: E402  同一份正文内容

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

OUT = os.path.join(
    M.OUT_DIR, "说明书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0.docx")

# pPr 里 pBdr 之后的合法子元素顺序（ECMA-376 CT_PPrBase）
PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")

SECTPR_AFTER_PGNUMTYPE = (
    "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
    "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
    "w:printerSettings", "w:sectPrChange")


def clear_theme_font(rpr):
    """去掉 rFonts 上的主题字体引用，否则 Word 会拿主题字体覆盖我们指定的中文字体。"""
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        return
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]


def force_black(style):
    """标题一律纯黑（Word 默认标题是主题蓝）。"""
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
            if color.get(qn(attr)) is not None:
                del color.attrib[qn(attr)]
    clear_theme_font(rpr)


def set_style_font(style, cn, en, size, bold=False):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    force_black(style)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:cs"), en)
    clear_theme_font(rpr)


def add_bottom_border(pr):
    """给段落或段落样式加下边框（PDF 版每章标题下、页眉下各有一条细线）。"""
    ppr = pr.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")          # 0.75pt
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "808080")
    pbdr.append(bottom)
    ppr.insert_element_before(pbdr, *PPR_AFTER_PBDR)


def add_page_field(paragraph):
    """插入「第 X 页」里的 PAGE 域（Word 打开后自动显示真实页码）。"""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def fit_size(path):
    """按正文宽度缩放，并按一页能容纳的高度封顶。

    与 PDF 版同一套缩放规则：长截图若只按宽度缩放，高度会超过一页，
    Word 会把它拦腰截到下一页去 —— 说明书里的界面截图被切开是硬伤。
    """
    info = DocxImage.from_file(path)
    nat_w = info.width * 72.0 / (info.horz_dpi or 72)
    nat_h = info.height * 72.0 / (info.vert_dpi or 72)
    avail = M.CONTENT_BOTTOM - M.CONTENT_TOP - 20      # 与 PDF 版一样给图注留 20pt
    scale = min(M.BODY_W / nat_w, avail / nat_h)
    return nat_w * scale, nat_h * scale


def page_setup(section):
    section.page_width = Pt(M.PAGE_W)
    section.page_height = Pt(M.PAGE_H)
    section.left_margin = Pt(M.MARGIN_L)
    section.right_margin = Pt(M.MARGIN_R)
    section.top_margin = Pt(M.MARGIN_T)
    section.bottom_margin = Pt(M.MARGIN_B)
    section.header_distance = Pt(M.HEADER_Y)
    section.footer_distance = Pt(30)


def restart_numbering(section):
    """页码自目录页重新从 1 起（封面不编号）。"""
    sect = section._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sect.insert_element_before(pg, *SECTPR_AFTER_PGNUMTYPE)


def build_header(section):
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Pt(M.BODY_W), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run("%s\t第 " % M.SOFTWARE_NAME)
    run.font.size = Pt(9)
    add_page_field(p)
    run2 = p.add_run(" 页")
    run2.font.size = Pt(9)
    add_bottom_border(p._p)


def build_cover(doc):
    t = doc.add_paragraph()
    t.paragraph_format.space_before = Pt(120)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("历史遗留废弃矿山自然恢复适宜性评价系统")
    r.font.size = Pt(22)
    r.font.name = CN_HEAD
    r._element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)

    v = doc.add_paragraph()
    v.alignment = WD_ALIGN_PARAGRAPH.CENTER
    v.paragraph_format.space_before = Pt(24)
    rv = v.add_run("V1.0")
    rv.font.size = Pt(16)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_before = Pt(72)
    rs = s.add_run("软件说明书（用户手册）")
    rs.font.size = Pt(18)
    rs.font.name = CN_HEAD
    rs._element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)

    first = doc.add_paragraph()
    first.paragraph_format.space_before = Pt(96)
    first.paragraph_format.left_indent = Pt(M.MARGIN_L + 60)
    for i, line in enumerate(("软件全称：历史遗留废弃矿山自然恢复适宜性评价系统",
                              "版本号：V1.0",
                              "文档类型：软件说明书（用户手册）")):
        p = first if i == 0 else doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(M.MARGIN_L + 60)
        p.paragraph_format.space_after = Pt(12)
        r0 = p.add_run(line)
        r0.font.size = Pt(11)


def add_bookmark(paragraph, name, bid):
    """在段落首尾插入书签，供目录内部跳转使用。

    注意书签必须插在 `w:pPr` **之后**：`w:pPr` 必须是段落的第一个子元素，
    插到它前面会让 Word 报「内容有问题」。
    """
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text, anchor, size):
    """把一段文字做成指向书签的内部超链接（外观仍是正文样式，不变成蓝色链接）。"""
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    paragraph._p.replace(run._r, link)
    link.append(run._r)


def build_toc(doc):
    """目录：每条都做成指向对应章节的内部链接，点一下就能跳过去。"""
    doc.add_heading("目　录", 1)
    for idx, (text, lvl) in enumerate(M.toc_entries()):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(lvl * 20)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(20)
        add_internal_link(p, text, "toc_%d" % idx, 11 if lvl == 0 else 10.5)


def build_body(doc, blocks):
    li_no = 0
    images = 0
    missing = []
    bookmark_no = 0
    for kind, *rest in blocks:
        if kind == "li":
            li_no += 1
            marker = M.cn_ordinal(li_no) + "、"
        else:
            li_no = 0

        if kind in ("h1", "h2"):
            h = doc.add_heading(M.plain(rest[0]), 1 if kind == "h1" else 2)
            add_bookmark(h, "toc_%d" % bookmark_no, 1000 + bookmark_no)
            bookmark_no += 1
        elif kind == "p":
            p = doc.add_paragraph(M.plain(rest[0]))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = Pt(17)     # 与 PDF 版正文行距一致
        elif kind == "li":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Pt(34)
            pf.first_line_indent = Pt(-34)      # 悬挂缩进，与 PDF 版一致
            pf.space_after = Pt(4)
            pf.line_spacing = Pt(17)
            p.add_run(marker + M.plain(rest[0]))
        elif kind == "img":
            name, caption = rest[0], rest[1]
            path = os.path.join(M.SHOT_DIR, name)
            if not os.path.exists(path):
                missing.append(name)
                continue
            w, h = fit_size(path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(10)
            pf.space_after = Pt(2)
            pf.keep_with_next = True
            p.add_run().add_picture(path, width=Pt(w), height=Pt(h))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(14)
            rc = cap.add_run(caption)
            rc.font.size = Pt(9)
            images += 1
    return images, missing


def main():
    os.makedirs(M.OUT_DIR, exist_ok=True)
    doc = Document()

    set_style_font(doc.styles["Normal"], CN_BODY, EN_FONT, 10.5)
    npf = doc.styles["Normal"].paragraph_format
    npf.space_before = Pt(0)
    npf.space_after = Pt(6)
    # 行距用「单倍」，具体段落的固定行距在正文块里逐段设置。
    # 注意：不能把 Normal 设成固定行距 —— 插图段落会继承它，
    # Word 认为该行只有 17pt 高，图片就会与后面的文字叠在一起（页数也会算错）。
    npf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    npf.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, CN_HEAD, EN_FONT, 14, bold=True)
    h1.paragraph_format.space_before = Pt(20)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    h1.paragraph_format.keep_with_next = True
    add_bottom_border(h1.element)

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, CN_HEAD, EN_FONT, 12, bold=True)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(12)
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    h2.paragraph_format.keep_with_next = True

    # 封面单独一节：不编页码、不加页眉
    page_setup(doc.sections[0])
    build_cover(doc)

    # 目录起另起一节，页码从 1 开始
    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(body_sec)
    build_header(body_sec)
    restart_numbering(body_sec)

    build_toc(doc)
    doc.add_page_break()
    images, missing = build_body(doc, M.CONTENT)

    doc.save(OUT)

    # 自检：正文里不许残留 Markdown 标记，也不许漏图
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    for bad, label in (("**", "Markdown 加粗标记"), ("`", "反引号"), ("·", "间隔点")):
        n = joined.count(bad)
        print("  %s：%d 处%s" % (label, n, "" if n == 0 else "  ← 需检查"))
    print("已生成 Word 版说明书：%s" % os.path.relpath(OUT, M.BASE))
    print("  段落 %d 段、插图 %d 张（正文块 %d 个）"
          % (len(texts), images, len(M.CONTENT)))
    print("  目录条目 %d 条（由正文标题推导）" % len(M.toc_entries()))
    print("  文件大小：%.2f MB" % (os.path.getsize(OUT) / 1024 / 1024))
    if missing:
        print("  [警告] 缺图：%s" % "、".join(missing))


if __name__ == "__main__":
    main()
