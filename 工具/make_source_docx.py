# -*- coding: utf-8 -*-
r"""生成软著源程序材料的 Word 版（与 PDF 版内容完全一致）。

为什么要 Word 版：代办/评审方要求把代码整理到 Word 里（附件3：尽量 3000 行以上）。
代码行直接复用 `make_source_doc.collect_lines()`，与 PDF 同源，保证两份材料一致。

排版与 PDF 对齐：
- A4 纵向；页眉左「软件名称+版本号」、右「第 N 页」（用 Word 页码域，自动更新）
- 每页 50 行，行号右对齐在右侧；字号 8pt、行距固定 14pt
- 西文用 Consolas（等宽），中文注释用宋体

输出：
    软著材料/源程序_提交版（前30页+后30页）_……V1.0.docx     前 30 页 + 后 30 页（3000 行）
    软著材料/源程序_……V1.0.docx                          全文
"""
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "工具"))
import make_source_doc as S          # noqa: E402  （复用同一份代码行收集逻辑）

OUT_DIR = ROOT / "软著材料"
LINES_PER_PAGE = S.LINES_PER_PAGE     # 50
CN = "宋体"
EN = "Consolas"
SIZE = Pt(S.FONT_SIZE)                # 8pt
LINE = Pt(S.LINE_HEIGHT)              # 14pt


def set_run(run, text, size=SIZE):
    run.text = text
    run.font.size = size
    run.font.name = EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    run._element.rPr.rFonts.set(qn("w:ascii"), EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), EN)


def setup_page(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin = Pt(S.MARGIN_L), Pt(S.MARGIN_R)
    s.top_margin, s.bottom_margin = Pt(52), Pt(48)
    s.header_distance = Pt(22)

    # 页眉：左软件名称，右页码（域）
    hp = s.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(
        s.page_width - s.left_margin - s.right_margin, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(S.SOFTWARE_NAME + "\t")
    r.font.size = Pt(10)
    r.font.name = CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    r2 = hp.add_run("第 ")
    r2.font.size = Pt(10)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    hp._p.append(fld)
    r3 = hp.add_run(" 页")
    r3.font.size = Pt(10)
    r3.font.name = CN
    r3._element.rPr.rFonts.set(qn("w:eastAsia"), CN)


def add_page(doc, chunk, first_of_doc):
    """写一页：50 行，行号在右侧（用右对齐制表位）。"""
    if not first_of_doc:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    s = doc.sections[0]
    right = s.page_width - s.left_margin - s.right_margin
    for ln, text in chunk:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = LINE
        pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
        set_run(p.add_run(), text + "\t")
        set_run(p.add_run(), str(ln))


def build_docx(all_lines, out_path, title):
    doc = Document()
    setup_page(doc)
    # 文档属性：便于代办识别
    doc.core_properties.title = title
    doc.core_properties.author = S.SOFTWARE_NAME
    for i in range(0, len(all_lines), LINES_PER_PAGE):
        add_page(doc, all_lines[i:i + LINES_PER_PAGE], first_of_doc=(i == 0))
    doc.save(str(out_path))
    return len(all_lines)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    print("收集源程序（与 PDF 同一份代码行）…")
    lines = S.collect_lines()
    numbered = [(i + 1, t) for i, t in enumerate(lines)]
    print("  有效代码行数：%d" % len(numbered))

    full = OUT_DIR / ("源程序_%s.docx" % S.SOFTWARE_NAME)
    n = build_docx(numbered, full, "源程序_全文_%s" % S.SOFTWARE_NAME)
    print("已生成全文 Word：%s（%d 行）" % (full.name, n))

    if len(numbered) > 60 * LINES_PER_PAGE:
        per = LINES_PER_PAGE
        sub_lines = numbered[:30 * per] + numbered[-30 * per:]
        sub = OUT_DIR / ("源程序_提交版（前30页+后30页）_%s.docx" % S.SOFTWARE_NAME)
        build_docx(sub_lines, sub, "源程序_提交版_%s" % S.SOFTWARE_NAME)
        print("已生成提交版 Word：%s（前 30 页 + 后 30 页 = %d 行，%d 页）"
              % (sub.name, len(sub_lines), len(sub_lines) // per))
    else:
        print("总行数不足 60 页，全文即提交版。")


if __name__ == "__main__":
    main()
