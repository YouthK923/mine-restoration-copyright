# -*- coding: utf-8 -*-
import os, sys, io, json

# 本脚本位于子目录，仓库根目录才是 app/ 与 软著材料/ 所在位置
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)
sys.stdout.reconfigure(encoding='utf-8')

BASE = ROOT
OUT = os.path.join(BASE, '参考资料', '论文与模板解析文本')
os.makedirs(OUT, exist_ok=True)

def dump_docx(path, outname):
    import docx
    d = docx.Document(path)
    lines = []
    # iterate body in document order (paragraphs + tables)
    from docx.oxml.ns import qn
    body = d.element.body
    pidx = 0
    tidx = 0
    paras = d.paragraphs
    tables = d.tables
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            if pidx < len(paras):
                p = paras[pidx]
                style = p.style.name if p.style is not None else ''
                txt = p.text
                if txt.strip() or style.startswith('Heading'):
                    prefix = f'[{style}] ' if style and style != 'Normal' else ''
                    lines.append(prefix + txt)
                pidx += 1
        elif child.tag == qn('w:tbl'):
            if tidx < len(tables):
                t = tables[tidx]
                lines.append('--- TABLE ---')
                for row in t.rows:
                    lines.append(' | '.join(c.text.replace('\n', ' ') for c in row.cells))
                lines.append('--- END TABLE ---')
                tidx += 1
    # inline images count
    imgs = [r for r in d.part.rels.values() if "image" in r.reltype]
    lines.append('')
    lines.append(f'### META: paragraphs={len(paras)} tables={len(tables)} images={len(imgs)}')
    with open(os.path.join(OUT, outname), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    return len(lines), len(imgs)

def dump_pdf(path, outname, maxpages=None):
    import fitz
    doc = fitz.open(path)
    lines = []
    lines.append(f'### PDF: {os.path.basename(path)} pages={doc.page_count}')
    for i, page in enumerate(doc):
        if maxpages and i >= maxpages:
            lines.append(f'... [truncated at page {maxpages}]')
            break
        lines.append(f'\n===== PAGE {i+1} =====')
        t = page.get_text('text')
        lines.append(t.rstrip())
        imgs = page.get_images(full=True)
        if imgs:
            lines.append(f'[page {i+1} images: {len(imgs)}]')
    with open(os.path.join(OUT, outname), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    return doc.page_count

res = {}
res['docx'] = dump_docx(os.path.join(BASE, '参考资料', '论文原稿.docx'), 'draft_docx.txt')
res['biansu_shuomingshu'] = dump_pdf(os.path.join(BASE, '参考资料', '同类软著样例/变压器抗短路能力评估系统V1.0-说明书.pdf'), 'tpl_biansu_manual.txt')
res['biansu_code'] = dump_pdf(os.path.join(BASE, '参考资料', '同类软著样例/变压器抗短路能力评估系统V1.0-代码.pdf'), 'tpl_biansu_code.txt', maxpages=3)
res['dianwang_manual'] = dump_pdf(os.path.join(BASE, '参考资料', '同类软著样例/电网安全风险数据分析系统-用户手册.pdf'), 'tpl_dianwang_manual.txt')
res['dianwang_code'] = dump_pdf(os.path.join(BASE, '参考资料', '同类软著样例/电网安全风险数据分析系统-源代码.pdf'), 'tpl_dianwang_code.txt', maxpages=6)
print(json.dumps(res, ensure_ascii=False, indent=1))
