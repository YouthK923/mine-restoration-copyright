# -*- coding: utf-8 -*-
import sys, re, zipfile
sys.stdout.reconfigure(encoding='utf-8')
import docx

FN = '软件著作权登记项目计划书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0.docx'
z = zipfile.ZipFile(FN)

print('=== 各 part 中残留的这些颜色 ===')
for n in z.namelist():
    if not n.endswith('.xml'):
        continue
    t = z.read(n).decode('utf-8', 'ignore')
    hits = {c: t.count(c) for c in ('4F81BD', '335B8A', '366091', '345A8A', 'themeColor')
            if c in t}
    if hits:
        print(' ', n, hits)

print()
print('=== styles.xml 里 4F81BD 所在样式 ===')
s = z.read('word/styles.xml').decode('utf-8')
for m in re.finditer(r'<w:style\b[^>]*?w:styleId="([^"]+)"[^>]*?>(.*?)</w:style>', s, re.S):
    if '4F81BD' in m.group(2):
        nm = re.search(r'<w:name w:val="([^"]*)"', m.group(2))
        print('  id=%s name=%s' % (m.group(1), nm.group(1) if nm else '?'))

print()
print('=== 改后各标题样式的颜色定义 ===')
for sid in ('2', '4', '5', '16', '39'):
    m = re.search(r'<w:style\b[^>]*?w:styleId="' + sid + r'"[^>]*?>(.*?)</w:style>', s, re.S)
    if m:
        nm = re.search(r'<w:name w:val="([^"]*)"', m.group(1))
        col = re.findall(r'<w:color\b[^>]*?/>', m.group(1))
        ea = re.findall(r'w:eastAsia="([^"]*)"', m.group(1))
        print('  id=%-3s %-14s color=%s  eastAsia=%s' % (sid, nm.group(1), col or '-', ea or '-'))

print()
print('=== 目录(TOC)段落检查 ===')
d = docx.Document(FN)
toc_like = [p for p in d.paragraphs
            if p.style is not None and ('TOC' in p.style.name or 'toc' in p.style.name.lower())]
print('  TOC 样式段落数:', len(toc_like))
for p in toc_like[:8]:
    print('   [%s] %s' % (p.style.name, p.text[:60]))
if not toc_like:
    # 目录可能是普通段落
    print('  未找到 TOC 样式段落，检查开头若干段落：')
    for p in d.paragraphs[:12]:
        if p.text.strip():
            print('   [%s] %s' % (p.style.name, p.text[:70]))
