# -*- coding: utf-8 -*-
"""Inspect heading styles + inline colors in the user's edited DOCX. Read-only."""
import sys, re, zipfile
sys.stdout.reconfigure(encoding='utf-8')
import docx

FN = '软件著作权登记项目计划书_历史遗留废弃矿山自然恢复适宜性评价系统V1.0.docx'
z = zipfile.ZipFile(FN)
s = z.read('word/styles.xml').decode('utf-8')

print('=== 1) 样式定义中的颜色 ===')
for m in re.finditer(r'<w:style [^>]*w:styleId="([^"]+)"[^>]*>(.*?)</w:style>', s, re.S):
    sid, blk = m.group(1), m.group(2)
    cols = re.findall(r'<w:color[^>]*w:val="([^"]+)"[^>]*/?>', blk)
    if cols:
        print('  %-14s color=%s' % (sid, cols))

print()
print('=== 2) 文档正文中的直接颜色标记 (inline w:color) ===')
doc = z.read('word/document.xml').decode('utf-8')
inline = re.findall(r'<w:color w:val="([^"]+)"', doc)
from collections import Counter
print('  ', dict(Counter(inline)))

print()
print('=== 3) 标题段落及其样式名 ===')
d = docx.Document(FN)
for p in d.paragraphs:
    st = p.style.name if p.style is not None else ''
    if st.startswith('Heading') or st == 'Title':
        # inline run colors
        rc = []
        for r in p.runs:
            c = r.font.color
            if c is not None and c.rgb is not None:
                rc.append(str(c.rgb))
        print('  [%-9s] %-46s inline_rgb=%s' % (st, p.text[:44], rc or '-'))

print()
print('=== 4) styles.xml 里 Heading/Title 完整块 ===')
for sid in ('Title', 'Heading1', 'Heading2', 'Heading3'):
    m = re.search(r'<w:style [^>]*w:styleId="' + sid + r'"[^>]*>.*?</w:style>', s, re.S)
    print('---', sid, '---')
    print(m.group(0) if m else 'NOT FOUND')

print()
print('=== 5) 文档基本信息 ===')
print('  段落数:', len(d.paragraphs), ' 表格数:', len(d.tables))
